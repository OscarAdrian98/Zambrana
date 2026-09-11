from dataclasses import dataclass
from datetime import date, timedelta
from typing import Any

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from app.services.horarios import (
    ReconstruccionHoraria,
    formatear_timedelta_hhmmss,
)


@dataclass(frozen=True)
class FilaExportacionPersonal:
    fecha: str
    registros: tuple[str, ...]
    total: timedelta
    total_texto: str
    incidencias: tuple[str, ...]

    @property
    def registros_texto(self) -> str:
        return "\n".join(self.registros) if self.registros else "-"

    @property
    def incidencias_texto(self) -> str:
        return "\n".join(self.incidencias)


def sanear_texto_excel(valor: Any) -> Any:
    """Fuerza como texto literal las cadenas que Excel interpretaría como fórmula."""
    if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
        return f"'{valor}"
    return valor


def _texto_seguro(valor: Any, sanear_excel: bool) -> str:
    texto = str(valor)
    if sanear_excel:
        texto = sanear_texto_excel(texto)
    return texto


def _origen(registro: Any, sanear_excel: bool) -> str:
    return _texto_seguro(getattr(registro, "origen", None) or "Tienda", sanear_excel)


def _texto_registro(registro: Any, sanear_excel: bool) -> str:
    tipo = _texto_seguro(
        str(getattr(registro, "tipo", "")).capitalize(),
        sanear_excel,
    )
    hora = registro.timestamp.strftime("%H:%M:%S")
    return f"{tipo} - {hora} ({_origen(registro, sanear_excel)})"


def _texto_tramo(tramo: Any, sanear_excel: bool) -> str:
    entrada = tramo.entrada.timestamp.strftime("%H:%M")
    salida = tramo.salida.timestamp.strftime("%H:%M")
    origen = _origen(tramo.entrada, sanear_excel)
    return f"{entrada}-{salida} ({origen})"


def _descripcion_incidencia(incidencia: Any, sanear_excel: bool) -> str:
    codigo = incidencia.codigo
    registro = incidencia.registro
    hora = registro.timestamp.strftime("%H:%M:%S")
    descripciones = {
        "doble_entrada": f"Doble entrada a las {hora}",
        "doble_salida": f"Doble salida a las {hora}",
        "salida_sin_entrada": f"Salida sin entrada a las {hora}",
        "duracion_no_positiva": f"Duración no positiva a las {hora}",
        "entrada_abierta": f"Entrada abierta a las {hora}",
    }
    if codigo == "tipo_desconocido":
        tipo = _texto_seguro(getattr(registro, "tipo", ""), sanear_excel)
        descripcion = f"Tipo desconocido ({tipo}) a las {hora}"
    else:
        descripcion = descripciones.get(codigo, f"Incidencia {codigo} a las {hora}")
    return _texto_seguro(descripcion, sanear_excel)


def preparar_incidencias(
    resultado: ReconstruccionHoraria,
    *,
    sanear_excel: bool = False,
) -> tuple[str, ...]:
    return tuple(
        _descripcion_incidencia(incidencia, sanear_excel)
        for incidencia in resultado.incidencias
    )


def preparar_fila_administrativa(
    nombre_usuario: str,
    fecha: date,
    resultado: ReconstruccionHoraria,
) -> dict[str, str]:
    primer_tramo = resultado.tramos[0] if resultado.tramos else None
    segundo_tramo = resultado.tramos[1] if len(resultado.tramos) > 1 else None
    adicionales = resultado.tramos[2:]

    return {
        "Usuario": sanear_texto_excel(nombre_usuario),
        "Fecha": fecha.strftime("%d/%m/%Y"),
        "Entrada Mañana": (
            primer_tramo.entrada.timestamp.strftime("%H:%M")
            if primer_tramo
            else ""
        ),
        "Salida Mañana": (
            primer_tramo.salida.timestamp.strftime("%H:%M") if primer_tramo else ""
        ),
        "Entrada Tarde": (
            segundo_tramo.entrada.timestamp.strftime("%H:%M")
            if segundo_tramo
            else ""
        ),
        "Salida Tarde": (
            segundo_tramo.salida.timestamp.strftime("%H:%M")
            if segundo_tramo
            else ""
        ),
        "Total trabajado": (
            formatear_timedelta_hhmmss(resultado.total_trabajado)
            if resultado.total_trabajado.total_seconds() > 0
            else ""
        ),
        "Tramos adicionales": "\n".join(
            _texto_tramo(tramo, True) for tramo in adicionales
        ),
        "Incidencias": "\n".join(
            preparar_incidencias(resultado, sanear_excel=True)
        ),
    }


def preparar_fila_personal(
    fecha: date,
    resultado: ReconstruccionHoraria,
    *,
    sanear_excel: bool = False,
) -> FilaExportacionPersonal:
    return FilaExportacionPersonal(
        fecha=fecha.strftime("%d/%m/%Y"),
        registros=tuple(
            _texto_registro(registro, sanear_excel)
            for registro in resultado.registros_activos
        ),
        total=resultado.total_trabajado,
        total_texto=(
            formatear_timedelta_hhmmss(resultado.total_trabajado)
            if resultado.total_trabajado.total_seconds() > 0
            else "-"
        ),
        incidencias=preparar_incidencias(
            resultado,
            sanear_excel=sanear_excel,
        ),
    )


def configurar_hoja_excel(
    worksheet,
    *,
    fila_cabecera: int,
    anchos: tuple[float, ...],
    columnas_multilinea: tuple[int, ...] = (),
    columnas_texto: tuple[int, ...] = (),
) -> None:
    ultima_columna = get_column_letter(len(anchos))
    worksheet.freeze_panes = f"A{fila_cabecera + 1}"
    worksheet.auto_filter.ref = (
        f"A{fila_cabecera}:{ultima_columna}{max(fila_cabecera, worksheet.max_row)}"
    )

    for indice, ancho in enumerate(anchos, start=1):
        worksheet.column_dimensions[get_column_letter(indice)].width = ancho

    for indice in columnas_multilinea:
        for fila in range(fila_cabecera + 1, worksheet.max_row + 1):
            worksheet.cell(fila, indice).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    for indice in columnas_texto:
        for fila in range(fila_cabecera + 1, worksheet.max_row + 1):
            worksheet.cell(fila, indice).number_format = "@"
