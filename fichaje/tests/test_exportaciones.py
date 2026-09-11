from datetime import date, datetime, time, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from PyPDF2 import PdfReader

from app import db
from app.models import Fichaje, RegistroHorario
from app.services.exportaciones import (
    preparar_fila_personal,
    sanear_texto_excel,
)
from app.services.horarios import reconstruir_tramos


def login(client, usuario):
    return client.post(
        "/login",
        data={
            "email": usuario.email,
            "password": "password-correcta",
        },
    )


def crear_fichaje(usuario, fecha, registros):
    fichaje = Fichaje(
        usuario_id=usuario.id,
        fecha=fecha,
        fecha_creacion=datetime.combine(fecha, time(7, 0)),
        creado_por_admin=False,
        eliminado=False,
    )
    db.session.add(fichaje)
    db.session.flush()

    for tipo, marca, origen, eliminado in registros:
        timestamp = (
            marca if isinstance(marca, datetime) else datetime.combine(fecha, marca)
        )
        db.session.add(
            RegistroHorario(
                fichaje_id=fichaje.id,
                tipo=tipo,
                timestamp=timestamp,
                creado_por_admin=False,
                eliminado=eliminado,
                origen=origen,
            )
        )

    db.session.commit()
    return fichaje


def registros_tramos(*tramos):
    registros = []
    for entrada, salida, origen in tramos:
        registros.extend(
            [
                ("entrada", entrada, origen, False),
                ("salida", salida, origen, False),
            ]
        )
    return registros


def libro_desde_respuesta(response):
    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return load_workbook(BytesIO(response.data))


def excel_administrativo(client, **filtros):
    response = client.get("/admin/exportar_excel", query_string=filtros)
    return libro_desde_respuesta(response)["Fichajes"]


def excel_personal(client, **filtros):
    response = client.get("/exportar_mis_fichajes", query_string=filtros)
    return libro_desde_respuesta(response)["Fichajes"]


def cabeceras(worksheet, fila):
    return {
        worksheet.cell(fila, columna).value: columna
        for columna in range(1, worksheet.max_column + 1)
    }


def fila_excel(worksheet, fila_cabecera, numero=1):
    columnas = cabeceras(worksheet, fila_cabecera)
    fila = fila_cabecera + numero
    return {
        nombre: worksheet.cell(fila, columna).value
        for nombre, columna in columnas.items()
    }


def texto_pdf(response):
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF")
    return "\n".join(
        pagina.extract_text() or ""
        for pagina in PdfReader(BytesIO(response.data)).pages
    )


def test_excel_admin_un_tramo(client, administrador, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos((time(9, 0), time(10, 30), "Tienda")),
    )
    login(client, administrador)

    worksheet = excel_administrativo(client)
    fila = fila_excel(worksheet, 1)

    assert fila["Entrada Mañana"] == "09:00"
    assert fila["Salida Mañana"] == "10:30"
    assert fila["Total trabajado"] == "01:30:00"
    assert fila["Incidencias"] is None


def test_excel_admin_dos_tramos(client, administrador, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos(
            (time(8, 0), time(9, 0), "Tienda"),
            (time(10, 0), time(12, 0), "Remoto"),
        ),
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Entrada Mañana"] == "08:00"
    assert fila["Salida Mañana"] == "09:00"
    assert fila["Entrada Tarde"] == "10:00"
    assert fila["Salida Tarde"] == "12:00"
    assert fila["Total trabajado"] == "03:00:00"


def test_excel_admin_tres_tramos_suma_total_y_muestra_adicional(
    client,
    administrador,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos(
            (time(8, 0), time(9, 0), "Tienda"),
            (time(10, 0), time(11, 0), "Tienda"),
            (time(15, 30), time(16, 15), "Remoto"),
        ),
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] == "02:45:00"
    assert fila["Tramos adicionales"] == "15:30-16:15 (Remoto)"


def test_excel_admin_cuatro_tramos_conserva_todos_los_adicionales(
    client,
    administrador,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos(
            (time(8, 0), time(9, 0), "Tienda"),
            (time(10, 0), time(11, 0), "Tienda"),
            (time(12, 0), time(13, 0), "Remoto"),
            (time(14, 0), time(16, 0), "Remoto"),
        ),
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] == "05:00:00"
    assert fila["Tramos adicionales"].splitlines() == [
        "12:00-13:00 (Remoto)",
        "14:00-16:00 (Remoto)",
    ]


def test_excel_admin_excluye_registro_eliminado(
    client,
    administrador,
    usuario,
):
    registros = registros_tramos((time(8, 0), time(9, 0), "Tienda"))
    registros.extend(
        [
            ("entrada", time(12, 34), "Oculto", True),
            ("salida", time(15, 34), "Oculto", True),
        ]
    )
    crear_fichaje(usuario, date.today(), registros)
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] == "01:00:00"
    assert "Oculto" not in str(fila["Tramos adicionales"])
    assert RegistroHorario.query.filter_by(eliminado=True).count() == 2


def test_excel_admin_entrada_abierta_no_se_suma_y_genera_incidencia(
    client,
    administrador,
    usuario,
):
    registros = registros_tramos((time(8, 0), time(9, 0), "Tienda"))
    registros.append(("entrada", time(10, 0), "Remoto", False))
    crear_fichaje(usuario, date.today(), registros)
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] == "01:00:00"
    assert "Entrada abierta" in fila["Incidencias"]


def test_excel_admin_doble_entrada_no_inventa_tiempo(
    client,
    administrador,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("entrada", time(9, 0), "Tienda", False),
            ("entrada", time(10, 0), "Remoto", False),
            ("salida", time(11, 0), "Tienda", False),
        ],
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] == "02:00:00"
    assert "Doble entrada" in fila["Incidencias"]


def test_excel_admin_salida_aislada_no_inventa_tiempo(
    client,
    administrador,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        [("salida", time(10, 0), "Tienda", False)],
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] is None
    assert "Salida sin entrada" in fila["Incidencias"]


def test_excel_admin_respeta_filtros_de_fecha(client, administrador, usuario):
    hoy = date.today()
    fechas = [hoy.replace(day=1), hoy.replace(day=2), hoy.replace(day=3)]
    for fecha in fechas:
        crear_fichaje(
            usuario,
            fecha,
            registros_tramos((time(8, 0), time(9, 0), "Tienda")),
        )
    login(client, administrador)

    filtro = fechas[1].isoformat()
    worksheet = excel_administrativo(client, desde=filtro, hasta=filtro)

    assert worksheet.max_row == 2
    assert worksheet["B2"].value == fechas[1].strftime("%d/%m/%Y")


def test_excel_admin_sanea_nombre_que_empieza_por_formula(
    client,
    administrador,
    usuario,
):
    usuario.nombre = "=2+2"
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos((time(8, 0), time(9, 0), "Tienda")),
    )
    db.session.commit()
    login(client, administrador)

    celda = excel_administrativo(client)["A2"]

    assert celda.value == "'=2+2"
    assert celda.data_type == "s"


@pytest.mark.parametrize("prefijo", ["+", "-", "@"])
def test_saneado_excel_trata_prefijos_peligrosos_como_texto(prefijo):
    valor = sanear_texto_excel(f"{prefijo}contenido")

    assert valor == f"'{prefijo}contenido"


def test_excel_admin_sanea_origen_en_tramo_adicional(
    client,
    administrador,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos(
            (time(8, 0), time(9, 0), "Tienda"),
            (time(10, 0), time(11, 0), "Tienda"),
            (time(12, 0), time(13, 0), "+Origen"),
        ),
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert "('+Origen)" in fila["Tramos adicionales"]


def test_excel_admin_total_superior_a_24_horas(
    client,
    administrador,
    usuario,
):
    fecha = date.today()
    crear_fichaje(
        usuario,
        fecha,
        [
            ("entrada", datetime.combine(fecha, time(8, 0)), "Tienda", False),
            (
                "salida",
                datetime.combine(fecha + timedelta(days=1), time(12, 0)),
                "Tienda",
                False,
            ),
        ],
    )
    login(client, administrador)

    fila = fila_excel(excel_administrativo(client), 1)

    assert fila["Total trabajado"] == "28:00:00"


def test_excel_admin_formato_de_hoja(client, administrador, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos((time(8, 0), time(9, 0), "Tienda")),
    )
    login(client, administrador)

    worksheet = excel_administrativo(client)
    headers = cabeceras(worksheet, 1)

    assert list(headers) == [
        "Usuario",
        "Fecha",
        "Entrada Mañana",
        "Salida Mañana",
        "Entrada Tarde",
        "Salida Tarde",
        "Total trabajado",
        "Tramos adicionales",
        "Incidencias",
    ]
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == f"A1:I{worksheet.max_row}"
    assert worksheet.cell(2, headers["Tramos adicionales"]).alignment.wrap_text
    assert worksheet.cell(2, headers["Incidencias"]).alignment.wrap_text


def test_excel_personal_muestra_registros_activos_ordenados_y_origen(
    client,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("salida", time(10, 0), "Remoto", False),
            ("entrada", time(9, 0), "Remoto", False),
        ],
    )
    login(client, usuario)

    fila = fila_excel(excel_personal(client), 3)

    assert fila["Registros"].splitlines() == [
        "Entrada - 09:00:00 (Remoto)",
        "Salida - 10:00:00 (Remoto)",
    ]


def test_excel_personal_ordena_registros_aunque_se_inserten_desordenados(
    client,
    usuario,
):
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("salida", time(12, 0), "Tienda", False),
            ("entrada", time(8, 0), "Tienda", False),
        ],
    )
    login(client, usuario)

    registros = fila_excel(excel_personal(client), 3)["Registros"].splitlines()

    assert registros[0].startswith("Entrada - 08:00:00")
    assert registros[1].startswith("Salida - 12:00:00")


def test_excel_personal_conserva_origen_de_cada_registro(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("entrada", time(9, 0), "Auto", False),
            ("salida", time(10, 0), "Remoto", False),
        ],
    )
    login(client, usuario)

    registros = fila_excel(excel_personal(client), 3)["Registros"]

    assert "Entrada - 09:00:00 (Auto)" in registros
    assert "Salida - 10:00:00 (Remoto)" in registros


def test_excel_personal_incluye_tres_tramos_completos(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos(
            (time(8, 0), time(9, 0), "Tienda"),
            (time(10, 0), time(11, 0), "Tienda"),
            (time(12, 0), time(14, 0), "Remoto"),
        ),
    )
    login(client, usuario)

    fila = fila_excel(excel_personal(client), 3)

    assert len(fila["Registros"].splitlines()) == 6
    assert fila["Total trabajado"] == "04:00:00"


def test_excel_personal_excluye_registro_eliminado(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("entrada", time(9, 0), "Tienda", False),
            ("salida", time(10, 0), "Tienda", False),
            ("entrada", time(12, 34), "Oculto", True),
        ],
    )
    login(client, usuario)

    fila = fila_excel(excel_personal(client), 3)

    assert "12:34:00" not in fila["Registros"]
    assert fila["Total trabajado"] == "01:00:00"


def test_excel_personal_total_diario_coincide_con_resumen(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos(
            (time(8, 0), time(10, 0), "Tienda"),
            (time(11, 0), time(14, 0), "Remoto"),
        ),
    )
    login(client, usuario)

    fila = fila_excel(excel_personal(client), 3)
    resumen = client.get("/resumen").get_data(as_text=True)

    assert fila["Total trabajado"] == "05:00:00"
    assert "05:00:00" in resumen


def test_excel_personal_total_acumulado_suma_varios_dias(client, usuario):
    hoy = date.today()
    crear_fichaje(
        usuario,
        hoy.replace(day=1),
        registros_tramos((time(8, 0), time(10, 0), "Tienda")),
    )
    crear_fichaje(
        usuario,
        hoy.replace(day=2),
        registros_tramos((time(8, 0), time(12, 0), "Tienda")),
    )
    login(client, usuario)

    worksheet = excel_personal(client)

    assert worksheet.cell(6, 1).value == "Total acumulado:"
    assert worksheet.cell(6, 3).value == "06:00:00"


def test_excel_personal_total_acumulado_superior_a_24_horas(client, usuario):
    hoy = date.today()
    for numero_dia in (1, 2, 3):
        crear_fichaje(
            usuario,
            hoy.replace(day=numero_dia),
            registros_tramos((time(8, 0), time(18, 0), "Tienda")),
        )
    login(client, usuario)

    worksheet = excel_personal(client)

    assert worksheet.cell(7, 3).value == "30:00:00"


def test_excel_personal_entrada_abierta_no_se_suma_pero_se_conserva(
    client,
    usuario,
):
    registros = registros_tramos((time(8, 0), time(9, 0), "Tienda"))
    registros.append(("entrada", time(10, 0), "Remoto", False))
    crear_fichaje(usuario, date.today(), registros)
    login(client, usuario)

    fila = fila_excel(excel_personal(client), 3)

    assert fila["Total trabajado"] == "01:00:00"
    assert "Entrada - 10:00:00 (Remoto)" in fila["Registros"]
    assert "Entrada abierta" in fila["Incidencias"]


def test_excel_personal_refleja_incidencias(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("entrada", time(9, 0), "Tienda", False),
            ("entrada", time(10, 0), "Remoto", False),
        ],
    )
    login(client, usuario)

    fila = fila_excel(excel_personal(client), 3)

    assert "Doble entrada" in fila["Incidencias"]
    assert "Entrada abierta" in fila["Incidencias"]


def test_excel_personal_conserva_saltos_y_activa_wrap_text(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos((time(9, 0), time(10, 0), "Tienda")),
    )
    login(client, usuario)

    worksheet = excel_personal(client)
    headers = cabeceras(worksheet, 3)
    celda = worksheet.cell(4, headers["Registros"])

    assert "\n" in celda.value
    assert celda.alignment.wrap_text is True
    assert worksheet.freeze_panes == "A4"
    assert worksheet.auto_filter.ref == "A3:D4"


def test_excel_personal_sanea_nombre_y_texto_de_registro(client, usuario):
    usuario.nombre = "=Empleado"
    crear_fichaje(
        usuario,
        date.today(),
        [
            ("=formula", time(9, 0), "+Origen", False),
        ],
    )
    db.session.commit()
    login(client, usuario)

    worksheet = excel_personal(client)
    fila = fila_excel(worksheet, 3)

    assert worksheet["A1"].value == "Horario de '=Empleado"
    assert worksheet["A1"].data_type == "s"
    assert fila["Registros"].startswith("'=formula")
    assert "('+Origen)" in fila["Registros"]


def test_excel_personal_respeta_filtros(client, usuario):
    hoy = date.today()
    fechas = [hoy.replace(day=1), hoy.replace(day=2), hoy.replace(day=3)]
    for fecha in fechas:
        crear_fichaje(
            usuario,
            fecha,
            registros_tramos((time(8, 0), time(9, 0), "Tienda")),
        )
    login(client, usuario)

    filtro = fechas[1].isoformat()
    worksheet = excel_personal(client, desde=filtro, hasta=filtro)
    fila = fila_excel(worksheet, 3)

    assert fila["Fecha"] == fechas[1].strftime("%d/%m/%Y")
    assert worksheet.cell(5, 1).value == "Total acumulado:"


@pytest.mark.parametrize(
    ("registros", "total", "texto_esperado"),
    [
        (
            registros_tramos((time(8, 0), time(9, 0), "Tienda")),
            "01:00:00",
            "Salida - 09:00:00",
        ),
        (
            registros_tramos(
                (time(8, 0), time(9, 0), "Tienda"),
                (time(10, 0), time(11, 0), "Tienda"),
                (time(12, 0), time(14, 0), "Remoto"),
            ),
            "04:00:00",
            "Salida - 14:00:00",
        ),
    ],
    ids=["un-tramo", "tres-tramos"],
)
def test_pdf_usa_filas_preparadas_para_uno_y_tres_tramos(
    client,
    usuario,
    registros,
    total,
    texto_esperado,
):
    fichaje = crear_fichaje(usuario, date.today(), registros)
    login(client, usuario)

    fila = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )
    texto_generado = texto_pdf(client.get("/exportar_mis_fichajes_pdf"))

    assert fila.total_texto == total
    assert texto_esperado in fila.registros_texto
    assert total in texto_generado
    assert texto_esperado in texto_generado


def test_pdf_preparacion_excluye_eliminados(client, usuario):
    fichaje = crear_fichaje(
        usuario,
        date.today(),
        [
            ("entrada", time(9, 0), "Tienda", False),
            ("salida", time(10, 0), "Tienda", False),
            ("entrada", time(12, 34), "Oculto", True),
        ],
    )
    login(client, usuario)

    fila = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )
    texto_generado = texto_pdf(client.get("/exportar_mis_fichajes_pdf"))

    assert "12:34:00" not in fila.registros_texto
    assert "12:34:00" not in texto_generado
    assert fila.total_texto == "01:00:00"


def test_pdf_preparacion_conserva_entrada_abierta_sin_sumarla(client, usuario):
    fichaje = crear_fichaje(
        usuario,
        date.today(),
        [
            ("entrada", time(8, 0), "Tienda", False),
            ("salida", time(9, 0), "Tienda", False),
            ("entrada", time(10, 0), "Remoto", False),
        ],
    )
    login(client, usuario)

    fila = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )
    texto_generado = texto_pdf(client.get("/exportar_mis_fichajes_pdf"))

    assert fila.total_texto == "01:00:00"
    assert "Entrada - 10:00:00 (Remoto)" in fila.registros_texto
    assert "Entrada abierta" in fila.incidencias_texto
    assert "01:00:00" in texto_generado
    assert "Entrada - 10:00:00 (Remoto)" in texto_generado
    assert "Entrada abierta" in texto_generado


def test_pdf_preparacion_total_superior_a_24_horas(client, usuario):
    fecha = date.today()
    fichaje = crear_fichaje(
        usuario,
        fecha,
        [
            ("entrada", datetime.combine(fecha, time(8, 0)), "Tienda", False),
            (
                "salida",
                datetime.combine(fecha + timedelta(days=1), time(12, 0)),
                "Tienda",
                False,
            ),
        ],
    )
    login(client, usuario)

    fila = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )
    texto_generado = texto_pdf(client.get("/exportar_mis_fichajes_pdf"))

    assert fila.total_texto == "28:00:00"
    assert "28:00:00" in texto_generado


def test_pdf_preparacion_incluye_incidencias(client, usuario):
    fichaje = crear_fichaje(
        usuario,
        date.today(),
        [("salida", time(10, 0), "Tienda", False)],
    )
    login(client, usuario)

    fila = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )
    texto_generado = texto_pdf(client.get("/exportar_mis_fichajes_pdf"))

    assert fila.total_texto == "-"
    assert "Salida sin entrada" in fila.incidencias_texto
    assert "Salida sin entrada" in texto_generado


def test_pdf_respeta_filtros_de_fecha(client, usuario):
    hoy = date.today()
    fechas = [hoy.replace(day=1), hoy.replace(day=2), hoy.replace(day=3)]
    for fecha in fechas:
        crear_fichaje(
            usuario,
            fecha,
            registros_tramos((time(8, 0), time(9, 0), "Tienda")),
        )
    login(client, usuario)

    filtro = fechas[1].isoformat()
    texto = texto_pdf(
        client.get(
            "/exportar_mis_fichajes_pdf",
            query_string={"desde": filtro, "hasta": filtro},
        )
    )

    assert fechas[1].strftime("%d/%m/%Y") in texto
    assert fechas[0].strftime("%d/%m/%Y") not in texto
    assert fechas[2].strftime("%d/%m/%Y") not in texto


def test_pdf_respuesta_real_conserva_nombre_y_total(client, usuario):
    crear_fichaje(
        usuario,
        date.today(),
        registros_tramos((time(8, 0), time(10, 0), "Tienda")),
    )
    login(client, usuario)

    response = client.get("/exportar_mis_fichajes_pdf")
    texto = texto_pdf(response)

    assert len(response.data) > 1000
    assert "attachment" in response.headers["Content-Disposition"]
    assert f"fichajes_{usuario.nombre.replace(' ', '_')}.pdf" in response.headers[
        "Content-Disposition"
    ]
    assert "02:00:00" in texto
    assert "Total acumulado" in texto


def test_coherencia_tres_tramos_entre_vistas_excel_y_pdf(
    client,
    administrador,
):
    fichaje = crear_fichaje(
        administrador,
        date.today(),
        registros_tramos(
            (time(8, 0), time(9, 0), "Tienda"),
            (time(10, 0), time(12, 0), "Tienda"),
            (time(13, 0), time(16, 0), "Remoto"),
        ),
    )
    login(client, administrador)

    resumen = client.get("/resumen").get_data(as_text=True)
    admin_web = client.get("/admin/fichajes").get_data(as_text=True)
    admin_excel = fila_excel(excel_administrativo(client), 1)
    personal_excel = fila_excel(excel_personal(client), 3)
    fila_pdf = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )

    assert "06:00:00" in resumen
    assert "06:00:00" in admin_web
    assert admin_excel["Total trabajado"] == "06:00:00"
    assert personal_excel["Total trabajado"] == "06:00:00"
    assert fila_pdf.total_texto == "06:00:00"


def test_coherencia_tramo_cerrado_mas_entrada_abierta(
    client,
    administrador,
):
    fichaje = crear_fichaje(
        administrador,
        date.today(),
        [
            ("entrada", time(8, 0), "Tienda", False),
            ("salida", time(9, 0), "Tienda", False),
            ("entrada", time(10, 0), "Remoto", False),
        ],
    )
    login(client, administrador)

    fichar = client.get("/fichar").get_data(as_text=True)
    resumen = client.get("/resumen").get_data(as_text=True)
    admin_excel = fila_excel(excel_administrativo(client), 1)
    personal_excel = fila_excel(excel_personal(client), 3)
    fila_pdf = preparar_fila_personal(
        fichaje.fecha,
        reconstruir_tramos(fichaje.registros),
    )

    assert 'id="contador-trabajo">01:00<' in fichar
    assert "const totalCerradoSegundos = 3600;" in fichar
    assert "01:00:00" in resumen
    assert admin_excel["Total trabajado"] == "01:00:00"
    assert personal_excel["Total trabajado"] == "01:00:00"
    assert "Entrada - 10:00:00 (Remoto)" in personal_excel["Registros"]
    assert fila_pdf.total_texto == "01:00:00"
    assert "Entrada - 10:00:00 (Remoto)" in fila_pdf.registros_texto
